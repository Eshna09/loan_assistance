"""
Data Service — port 8003

Owns the knowledge-base documents and nothing else. It is the only component
that touches disk: it loads the bundled loan documents, accepts uploads,
extracts text from .txt/.md/.pdf/.docx, and serves chunks to the Retrieval
Service.

It deliberately has no embedding model and no FAISS index, which keeps its
container small and its responsibility single.

Endpoints
    GET    /health              liveness
    GET    /version             index-invalidation counter
    GET    /documents           document metadata for the dashboard
    GET    /chunks              every chunk, in stable order, for indexing
    POST   /documents           multipart upload of one or more files
    DELETE /documents/{name}    remove an uploaded document
"""

import io
import os
import re
import threading
from typing import List, Optional

from fastapi import FastAPI, File, Form, HTTPException, UploadFile

from services.common.chunking import CHUNK_SIZE, OVERLAP, chunk_text
from services.data_service.metadata_store import (
    delete_metadata,
    load_metadata,
    make_upload_metadata,
    save_metadata,
)

SERVICE_NAME = "data-service"

# Bundled documents ship with the repo; uploads live on a writable volume.
KB_DIR = os.environ.get(
    "KB_DIR",
    os.path.join(os.path.dirname(__file__), "..", "..", "backend", "finance_kb"),
)
UPLOAD_DIR = os.environ.get(
    "UPLOAD_DIR",
    os.path.join(os.path.dirname(__file__), "..", "..", "backend", "uploaded_kb"),
)

SUPPORTED_EXTENSIONS = {".txt", ".md", ".pdf", ".docx"}
MAX_UPLOAD_BYTES = 5 * 1024 * 1024  # 5 MB per file

BASE_DOCUMENT_FILENAMES = [
    "loans.txt",
    "loan_types.txt",
    "loan_interest.txt",
    "loan_repayment.txt",
    "loan_eligibility.txt",
    "loan_risks.txt",
]

os.makedirs(UPLOAD_DIR, exist_ok=True)

app = FastAPI(title="Data Service", version="1.0.0")

_lock = threading.RLock()
_version = 0  # bumped on every mutation so Retrieval knows to re-index


class DocumentError(ValueError):
    """Raised when an uploaded document cannot be accepted."""


# ── Text extraction ──────────────────────────────────────────────────────────
def _extract_pdf(raw: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(raw))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def _extract_docx(raw: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(raw))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def extract_text(filename: str, raw: bytes) -> str:
    """Turn an uploaded file's bytes into plain text, normalised for chunking."""
    ext = os.path.splitext(filename)[1].lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise DocumentError(
            f"Unsupported file type '{ext or filename}'. "
            f"Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )

    if ext == ".pdf":
        text = _extract_pdf(raw)
    elif ext == ".docx":
        text = _extract_docx(raw)
    else:
        text = raw.decode("utf-8", errors="replace")

    # Collapse the whitespace that PDF/DOCX extraction leaves behind so 300-char
    # chunks carry real content rather than blank runs.
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _safe_filename(filename: str) -> str:
    """Strip path components and characters that could escape UPLOAD_DIR."""
    base = os.path.basename(filename.replace("\\", "/")).strip()
    base = re.sub(r"[^A-Za-z0-9._ -]", "_", base)
    base = base.lstrip(".") or "document"
    root, ext = os.path.splitext(base)
    return (root[:80] or "document") + ext.lower()


def _unique_upload_name(filename: str) -> str:
    root, ext = os.path.splitext(filename)
    candidate = filename
    counter = 1
    while (
        os.path.exists(os.path.join(UPLOAD_DIR, candidate))
        or candidate in BASE_DOCUMENT_FILENAMES
    ):
        candidate = f"{root}_{counter}{ext}"
        counter += 1
    return candidate


# ── Document loading ─────────────────────────────────────────────────────────
def _uploaded_files_on_disk() -> list[str]:
    if not os.path.isdir(UPLOAD_DIR):
        return []
    return sorted(
        name
        for name in os.listdir(UPLOAD_DIR)
        if os.path.splitext(name)[1].lower() in SUPPORTED_EXTENSIONS
        and os.path.isfile(os.path.join(UPLOAD_DIR, name))
    )


def _read_document(filename: str) -> str:
    upload_path = os.path.join(UPLOAD_DIR, filename)
    if os.path.exists(upload_path):
        ext = os.path.splitext(filename)[1].lower()
        if ext in (".pdf", ".docx"):
            with open(upload_path, "rb") as f:
                return extract_text(filename, f.read())
        with open(upload_path, "r", encoding="utf-8", errors="replace") as f:
            return f.read()

    with open(os.path.join(KB_DIR, filename), "r", encoding="utf-8") as f:
        return f.read()


def _load_all() -> tuple[list[str], dict[str, str], set[str]]:
    """Return (ordered filenames, filename -> text, uploaded filenames)."""
    uploads = _uploaded_files_on_disk()
    filenames = BASE_DOCUMENT_FILENAMES + uploads
    raw = {name: _read_document(name) for name in filenames}
    return filenames, raw, set(uploads)


# ── Endpoints ────────────────────────────────────────────────────────────────
@app.get("/")
def root():
    return {"service": SERVICE_NAME, "message": "Data Service is running."}


@app.get("/health")
def health():
    with _lock:
        filenames, _, uploads = _load_all()
    return {
        "service": SERVICE_NAME,
        "status": "ok",
        "documents": len(filenames),
        "uploaded": len(uploads),
        "version": _version,
    }


@app.get("/version")
def version():
    """Cheap poll used by the Retrieval Service to detect a stale index."""
    return {"version": _version}


@app.get("/documents")
def documents():
    """Document metadata — what the dashboard renders in the Knowledge Base panel."""
    with _lock:
        filenames, raw, uploads = _load_all()
        docs = []
        total_chunks = 0
        for name in filenames:
            text = raw[name]
            count = len(chunk_text(text))
            total_chunks += count
            meta = load_metadata(UPLOAD_DIR, name)
            docs.append({
                "filename": name,
                "preview": text[:120],
                "chunk_count": count,
                "full_text": text,
                "uploaded": name in uploads,
                # metadata fields (additive)
                "document_type": meta.get("document_type", "general"),
                "loan_type": meta.get("loan_type", "general"),
                "version": meta.get("version", "1.0"),
                "effective_date": meta.get("effective_date"),
                "upload_date": meta.get("upload_date"),
            })
        return {
            "documents": docs,
            "total_chunks": total_chunks,
            "uploaded_count": len(uploads),
            "chunk_size": CHUNK_SIZE,
            "overlap": OVERLAP,
            "supported_upload_types": sorted(SUPPORTED_EXTENSIONS),
            "max_upload_mb": MAX_UPLOAD_BYTES // 1_048_576,
            "version": _version,
        }


@app.get("/chunks")
def chunks():
    """
    Every chunk in a stable order, with its source document.

    The Retrieval Service embeds this list positionally, so index i here is
    FAISS row i there.
    """
    with _lock:
        filenames, raw, _ = _load_all()
        items = []
        # Cache metadata per file to avoid re-loading on every chunk
        meta_cache: dict[str, dict] = {}
        for name in filenames:
            if name not in meta_cache:
                meta_cache[name] = load_metadata(UPLOAD_DIR, name)
            meta = meta_cache[name]
            for chunk in chunk_text(raw[name]):
                items.append({
                    "id": len(items),
                    "text": chunk,
                    "source": name,
                    # metadata fields (additive)
                    "document_type": meta.get("document_type", "general"),
                    "loan_type": meta.get("loan_type", "general"),
                    "version": meta.get("version", "1.0"),
                    "effective_date": meta.get("effective_date"),
                })
        return {"chunks": items, "total": len(items), "version": _version}


@app.post("/documents")
async def upload_documents(
    files: List[UploadFile] = File(...),
    document_type: Optional[str] = Form(None),
    loan_type: Optional[str] = Form(None),
    version: Optional[str] = Form(None),
    effective_date: Optional[str] = Form(None),
):
    """Accept one or more documents. Files are processed independently.
    
    Optional metadata Form fields (all backward-compatible):
    - document_type: e.g. "reference", "policy", "general"
    - loan_type: e.g. "home", "personal", "general"
    - version: e.g. "1.0", "2.1"
    - effective_date: ISO date string, e.g. "2024-01-01"
    """
    global _version

    if not files:
        raise HTTPException(status_code=400, detail="No files were uploaded.")

    uploaded = []
    failed = []

    for upload in files:
        try:
            raw = await upload.read()
            if not raw:
                raise DocumentError(f"'{upload.filename}' is empty.")
            if len(raw) > MAX_UPLOAD_BYTES:
                raise DocumentError(
                    f"'{upload.filename}' is {len(raw) / 1_048_576:.1f} MB — the limit is "
                    f"{MAX_UPLOAD_BYTES // 1_048_576} MB."
                )

            safe_name = _safe_filename(upload.filename)
            text = extract_text(safe_name, raw)
            if not text.strip():
                raise DocumentError(
                    f"No readable text found in '{upload.filename}'. "
                    "Scanned or image-only PDFs are not supported."
                )

            with _lock:
                # Persisted as .txt so a restart never re-runs extraction.
                stored = _unique_upload_name(os.path.splitext(safe_name)[0] + ".txt")
                with open(os.path.join(UPLOAD_DIR, stored), "w", encoding="utf-8") as f:
                    f.write(text)
                # Save metadata sidecar
                meta = make_upload_metadata(stored, document_type, loan_type, version, effective_date)
                save_metadata(UPLOAD_DIR, stored, meta)
                _version += 1

            uploaded.append({
                "filename": stored,
                "original_filename": upload.filename,
                "chunk_count": len(chunk_text(text)),
                "characters": len(text),
                "preview": text[:120],
                # return metadata so callers know what was stored
                "document_type": meta["document_type"],
                "loan_type": meta["loan_type"],
                "version": meta["version"],
                "effective_date": meta["effective_date"],
            })
        except DocumentError as exc:
            failed.append({"filename": upload.filename, "error": str(exc)})
        except Exception as exc:  # extraction failures from pypdf / python-docx
            failed.append({
                "filename": upload.filename,
                "error": f"Could not read the file: {exc}",
            })
        finally:
            await upload.close()

    if not uploaded and failed:
        raise HTTPException(status_code=400, detail=failed[0]["error"])

    return {"uploaded": uploaded, "failed": failed, "version": _version}


@app.delete("/documents/{filename}")
def delete_document(filename: str):
    """Remove an uploaded document. Bundled documents are protected."""
    global _version

    safe_name = _safe_filename(filename)
    with _lock:
        if safe_name in BASE_DOCUMENT_FILENAMES:
            raise HTTPException(
                status_code=400,
                detail=f"'{safe_name}' is a built-in knowledge-base document and cannot be deleted.",
            )
        path = os.path.join(UPLOAD_DIR, safe_name)
        if not os.path.isfile(path):
            raise HTTPException(
                status_code=404, detail=f"'{safe_name}' is not an uploaded document."
            )
        os.remove(path)
        delete_metadata(UPLOAD_DIR, safe_name)
        _version += 1

    return {"deleted": safe_name, "version": _version}
