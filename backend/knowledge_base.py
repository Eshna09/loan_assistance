"""
knowledge_base.py
Handles document loading, chunking, embedding, and FAISS indexing.
Core chunk/embed/retrieve logic is taken directly from the eshnaAIDevOps notebook.

Two document sources feed the index:
  * finance_kb/     — the bundled loan documents that ship with the project
  * uploaded_kb/    — documents uploaded at runtime through POST /kb/upload

Any change to either source triggers rebuild_index(), which re-chunks and
re-embeds everything so chunk_metadata stays aligned with the FAISS row order.
"""

import os
import io
import re
import threading

import numpy as np
from sentence_transformers import SentenceTransformer
import faiss

# ── Constants (from notebook) ───────────────────────────────────────────────
CHUNK_SIZE = 300
OVERLAP = 50
EMBEDDING_MODEL_NAME = "sentence-transformers/all-MiniLM-L6-v2"
EMBEDDING_DIMENSION = 384
KB_DIR = os.path.join(os.path.dirname(__file__), "finance_kb")
UPLOAD_DIR = os.path.join(os.path.dirname(__file__), "uploaded_kb")
FAISS_INDEX_TYPE = "IndexFlatL2"

# Upload limits
SUPPORTED_EXTENSIONS = {".txt", ".md", ".pdf", ".docx"}
MAX_UPLOAD_BYTES = 5 * 1024 * 1024  # 5 MB per file

# Bundled documents, in a fixed order so the UI list is stable.
BASE_DOCUMENT_FILENAMES = [
    "loans.txt",
    "loan_types.txt",
    "loan_interest.txt",
    "loan_repayment.txt",
    "loan_eligibility.txt",
    "loan_risks.txt",
]

os.makedirs(UPLOAD_DIR, exist_ok=True)


class DocumentError(ValueError):
    """Raised when an uploaded document cannot be accepted."""


# ── Chunking (exact algorithm from notebook) ─────────────────────────────────
def chunk_text(text: str, chunk_size: int = CHUNK_SIZE, overlap: int = OVERLAP) -> list[str]:
    """Sliding-window chunker matching the notebook implementation."""
    chunks = []
    start = 0
    while start < len(text):
        end = start + chunk_size
        chunks.append(text[start:end])
        start += (chunk_size - overlap)
    return chunks


# ── Text extraction for uploaded files ───────────────────────────────────────
def _extract_pdf(raw: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(raw))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def _extract_docx(raw: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(raw))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def extract_text(filename: str, raw: bytes) -> str:
    """Turn an uploaded file's bytes into plain text, normalised for chunking."""
    ext = os.path.splitext(filename)[1].lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise DocumentError(
            f"Unsupported file type '{ext or filename}'. "
            f"Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )

    if ext == ".pdf":
        text = _extract_pdf(raw)
    elif ext == ".docx":
        text = _extract_docx(raw)
    else:
        text = raw.decode("utf-8", errors="replace")

    # Collapse the whitespace that PDF/DOCX extraction tends to leave behind so
    # 300-char chunks carry real content rather than blank runs.
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _safe_filename(filename: str) -> str:
    """Strip any path components and characters that could escape UPLOAD_DIR."""
    base = os.path.basename(filename.replace("\\", "/")).strip()
    base = re.sub(r"[^A-Za-z0-9._ -]", "_", base)
    base = base.lstrip(".") or "document"
    root, ext = os.path.splitext(base)
    return (root[:80] or "document") + ext.lower()


def _unique_upload_name(filename: str) -> str:
    """Avoid clobbering an existing upload or shadowing a bundled document."""
    root, ext = os.path.splitext(filename)
    candidate = filename
    counter = 1
    while (
        os.path.exists(os.path.join(UPLOAD_DIR, candidate))
        or candidate in BASE_DOCUMENT_FILENAMES
    ):
        candidate = f"{root}_{counter}{ext}"
        counter += 1
    return candidate


# ── Module-level index state ─────────────────────────────────────────────────
print("Loading embedding model (MiniLM)...")
embedding_model = SentenceTransformer(EMBEDDING_MODEL_NAME)

_lock = threading.RLock()

document_filenames: list[str] = []
documents_raw: dict[str, str] = {}
uploaded_filenames: set[str] = set()
all_chunks: list[str] = []
chunk_metadata: list[dict] = []
faiss_index = faiss.IndexFlatL2(EMBEDDING_DIMENSION)


def _uploaded_files_on_disk() -> list[str]:
    if not os.path.isdir(UPLOAD_DIR):
        return []
    names = [
        name
        for name in os.listdir(UPLOAD_DIR)
        if os.path.splitext(name)[1].lower() in SUPPORTED_EXTENSIONS
        and os.path.isfile(os.path.join(UPLOAD_DIR, name))
    ]
    return sorted(names)


def _read_document(filename: str) -> str:
    """Read one document from whichever directory holds it."""
    upload_path = os.path.join(UPLOAD_DIR, filename)
    if os.path.exists(upload_path):
        ext = os.path.splitext(filename)[1].lower()
        if ext in (".pdf", ".docx"):
            with open(upload_path, "rb") as f:
                return extract_text(filename, f.read())
        with open(upload_path, "r", encoding="utf-8", errors="replace") as f:
            return f.read()

    with open(os.path.join(KB_DIR, filename), "r", encoding="utf-8") as f:
        return f.read()


def rebuild_index(verbose: bool = True) -> None:
    """Re-chunk, re-embed, and rebuild the FAISS index from all current documents."""
    global document_filenames, documents_raw, uploaded_filenames
    global all_chunks, chunk_metadata, faiss_index

    with _lock:
        uploads = _uploaded_files_on_disk()
        filenames = BASE_DOCUMENT_FILENAMES + uploads

        raw: dict[str, str] = {}
        chunks: list[str] = []
        metadata: list[dict] = []

        for filename in filenames:
            text = _read_document(filename)
            raw[filename] = text
            for chunk in chunk_text(text):
                chunks.append(chunk)
                metadata.append({"source": filename, "text": chunk})

        if verbose:
            print(f"Total chunks: {len(chunks)}")
            print("Converting text to vectors...")

        index = faiss.IndexFlatL2(EMBEDDING_DIMENSION)
        if chunks:
            embeddings = embedding_model.encode(chunks)
            index.add(np.array(embeddings, dtype=np.float32))

        document_filenames = filenames
        documents_raw = raw
        uploaded_filenames = set(uploads)
        all_chunks = chunks
        chunk_metadata = metadata
        faiss_index = index

        if verbose:
            print(f"Embedding dimension: {EMBEDDING_DIMENSION}")
            print(f"Successfully stored {faiss_index.ntotal} vectors in FAISS.")


rebuild_index()


# ── Upload / delete ──────────────────────────────────────────────────────────
def add_document(filename: str, raw: bytes) -> dict:
    """Validate, extract, persist, and index one uploaded document."""
    if len(raw) > MAX_UPLOAD_BYTES:
        raise DocumentError(
            f"'{filename}' is {len(raw) / 1_048_576:.1f} MB — the limit is "
            f"{MAX_UPLOAD_BYTES // 1_048_576} MB."
        )

    safe_name = _safe_filename(filename)
    text = extract_text(safe_name, raw)  # raises DocumentError on an unsupported type

    if not text.strip():
        raise DocumentError(
            f"No readable text found in '{filename}'. "
            "Scanned or image-only PDFs are not supported."
        )

    with _lock:
        # Everything is persisted as .txt so a restart never re-runs extraction.
        stored_name = _unique_upload_name(os.path.splitext(safe_name)[0] + ".txt")
        with open(os.path.join(UPLOAD_DIR, stored_name), "w", encoding="utf-8") as f:
            f.write(text)

        rebuild_index(verbose=False)
        chunk_count = len(chunk_text(text))

    return {
        "filename": stored_name,
        "original_filename": filename,
        "chunk_count": chunk_count,
        "characters": len(text),
        "preview": text[:120],
    }


def remove_document(filename: str) -> None:
    """Delete an uploaded document and rebuild the index. Bundled docs are protected."""
    safe_name = _safe_filename(filename)

    with _lock:
        if safe_name in BASE_DOCUMENT_FILENAMES:
            raise DocumentError(
                f"'{safe_name}' is a built-in knowledge-base document and cannot be deleted."
            )

        path = os.path.join(UPLOAD_DIR, safe_name)
        if not os.path.isfile(path):
            raise DocumentError(f"'{safe_name}' is not an uploaded document.")

        os.remove(path)
        rebuild_index(verbose=False)


# ── Retrieval (exact logic from notebook) ─────────────────────────────────────
def retrieve_context(question: str, top_k: int = 3):
    """Convert question to vector, search FAISS, return chunks + sources + distances."""
    with _lock:
        if faiss_index.ntotal == 0:
            return [], [], []

        question_embedding = embedding_model.encode([question])
        k = min(top_k, faiss_index.ntotal)
        distances, indices = faiss_index.search(np.array(question_embedding, dtype=np.float32), k)

        retrieved_chunks = []
        sources = []
        raw_distances = []
        for i, idx in enumerate(indices[0]):
            if idx < 0:
                continue
            chunk_data = chunk_metadata[idx]
            retrieved_chunks.append({"text": chunk_data["text"], "source": chunk_data["source"]})
            sources.append(chunk_data["source"])
            raw_distances.append(float(distances[0][i]))

        return retrieved_chunks, list(dict.fromkeys(sources)), raw_distances


# ── KB metadata for /kb/info ─────────────────────────────────────────────────
def get_kb_info() -> dict:
    """Return knowledge-base metadata for the dashboard."""
    with _lock:
        docs = []
        for filename in document_filenames:
            text = documents_raw[filename]
            docs.append({
                "filename": filename,
                "preview": text[:120],
                "chunk_count": len(chunk_text(text)),
                "full_text": text,
                "uploaded": filename in uploaded_filenames,
            })
        return {
            "documents": docs,
            "total_chunks": faiss_index.ntotal,
            "embedding_dimension": EMBEDDING_DIMENSION,
            "faiss_index_type": FAISS_INDEX_TYPE,
            "uploaded_count": len(uploaded_filenames),
            "supported_upload_types": sorted(SUPPORTED_EXTENSIONS),
            "max_upload_mb": MAX_UPLOAD_BYTES // 1_048_576,
        }


# ── Query embedding preview helper ───────────────────────────────────────────
def get_query_embedding_preview(question: str) -> list[float]:
    """Return first 8 + last 2 floats of the 384-dim query embedding."""
    vec = embedding_model.encode([question])[0]
    preview = list(vec[:8]) + list(vec[-2:])
    return [round(float(v), 4) for v in preview]
